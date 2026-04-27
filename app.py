import io
import re

import streamlit as st
import streamlit.components.v1 as components
import tiktoken
from pypdf import PdfReader
from docx import Document
from openai import OpenAI
from rdflib import Graph, RDF, OWL
from poml import poml
from pyvis.network import Network
from loguru import logger

from config import (
    OPENAI_API_KEY,
    KNOWN_MODEL_LIMITS,
    TIKTOKEN_ENCODING,
    BATCH_TOKEN_BUDGET,
    CHUNK_TOKEN_LIMIT,
    ACCEPTED_FILE_TYPES,
    PROMPTS_DIR,
)
from graph import (
    Neo4jClient,
    ontology_to_schema,
    build_knowledge_graph,
    query_graph_rag,
    query_agentic_rag,
    get_graph_stats,
    find_duplicate_entities,
    has_any_entities,
    compute_edge_weights,
    embed_all_entity_types,
    embed_relationships,
    embed_entities_by_type,
    get_embedding_stats,
    ENTITY_INDEX_CONFIGS,
    EMBEDDING_DIMENSIONS,
)
from entity_resolution import (
    find_candidate_pairs,
    score_exact,
    score_embedding_batch,
    score_llm_batch,
    build_transitive_clusters,
    merge_cluster,
    undo_merge,
    find_orphaned_nodes,
)
from web_sources import (
    extract_topics,
    search_and_fetch,
    build_web_knowledge_graph,
    compute_similar_to_edges,
    remove_web_content,
    get_web_source_stats,
)

ENCODING = tiktoken.get_encoding(TIKTOKEN_ENCODING)


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------
def extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()

    if name.endswith(".txt"):
        # Try multiple encodings to handle various file encodings
        raw_bytes = uploaded_file.read()
        for encoding in ("utf-8", "latin-1", "cp1252", "iso-8859-1", "utf-16"):
            try:
                return raw_bytes.decode(encoding)
            except (UnicodeDecodeError, UnicodeError):
                continue
        # Fallback: replace problematic characters
        return raw_bytes.decode("utf-8", errors="replace")

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(uploaded_file.read()))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(uploaded_file.read()))
        return "\n".join(para.text for para in doc.paragraphs)

    return ""


# ---------------------------------------------------------------------------
# Token helpers
# ---------------------------------------------------------------------------
def count_tokens(text: str) -> int:
    return len(ENCODING.encode(text))


# ---------------------------------------------------------------------------
# Chunking — general-purpose paragraph splitting
# ---------------------------------------------------------------------------
def chunk_text(text: str) -> list[str]:
    paragraphs = re.split(r"\n\n+", text.strip())
    paragraphs = [p.strip() for p in paragraphs if p.strip()]

    chunks: list[str] = []
    current_parts: list[str] = []
    current_tokens = 0

    for para in paragraphs:
        para_tokens = count_tokens(para)

        if para_tokens > CHUNK_TOKEN_LIMIT:
            if current_parts:
                chunks.append("\n\n".join(current_parts))
                current_parts = []
                current_tokens = 0
            chunks.append(para)
            logger.warning(f"Single paragraph exceeds chunk limit ({para_tokens} tokens)")
            continue

        if current_tokens + para_tokens > CHUNK_TOKEN_LIMIT and current_parts:
            chunks.append("\n\n".join(current_parts))
            current_parts = []
            current_tokens = 0

        current_parts.append(para)
        current_tokens += para_tokens

    if current_parts:
        chunks.append("\n\n".join(current_parts))

    return chunks


def chunk_all_documents(docs: list[dict]) -> list[str]:
    all_chunks: list[str] = []
    for doc in docs:
        file_chunks = chunk_text(doc["text"])
        all_chunks.extend(file_chunks)
    return all_chunks


# ---------------------------------------------------------------------------
# Batching — greedy bin-packing
# ---------------------------------------------------------------------------
def batch_chunks(chunks: list[str]) -> list[list[str]]:
    batches: list[list[str]] = []
    current_batch: list[str] = []
    current_tokens = 0

    for chunk in chunks:
        chunk_tokens = count_tokens(chunk)
        if current_tokens + chunk_tokens > BATCH_TOKEN_BUDGET and current_batch:
            batches.append(current_batch)
            current_batch = []
            current_tokens = 0
        current_batch.append(chunk)
        current_tokens += chunk_tokens

    if current_batch:
        batches.append(current_batch)

    return batches


# ---------------------------------------------------------------------------
# TTL helpers
# ---------------------------------------------------------------------------
def strip_code_fences(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(?:turtle|ttl|sparql)?\s*\n", "", text)
    text = re.sub(r"\n```\s*$", "", text)
    return text.strip()


def validate_ttl(ttl_string: str) -> tuple[bool, str]:
    try:
        g = Graph()
        g.parse(data=ttl_string, format="turtle")
        return True, g.serialize(format="turtle")
    except Exception as e:
        return False, str(e)


def merge_ttl_fragments(fragments: list[str]) -> str:
    merged = Graph()
    for frag in fragments:
        try:
            merged.parse(data=frag, format="turtle")
        except Exception as e:
            logger.warning(f"Skipping invalid fragment during merge: {e}")
    return merged.serialize(format="turtle")


def extract_registry(ttl_string: str) -> set[str]:
    g = Graph()
    g.parse(data=ttl_string, format="turtle")
    names: set[str] = set()
    for rdf_type in (OWL.Class, OWL.ObjectProperty, OWL.DatatypeProperty):
        for s in g.subjects(RDF.type, rdf_type):
            local_name = str(s).split("#")[-1].split("/")[-1]
            if local_name:
                names.add(local_name)
    return names


# ---------------------------------------------------------------------------
# LLM calls
# ---------------------------------------------------------------------------
def call_prompt_1a(client: OpenAI, model: str, doc_text: str, registry_text: str = "") -> str:
    # Read poml file with UTF 8 encoding
    poml_path = PROMPTS_DIR / "ontology_prompt.poml"
    with open(poml_path, "r", encoding="utf-8", errors="replace") as f:
        poml_content = f.read()
    
    params = poml(
        poml_content,
        context={"documents": doc_text, "registry": registry_text},
        format="openai_chat",
    )
    params["model"] = model

    response = client.chat.completions.create(**params)
    raw = response.choices[0].message.content
    return strip_code_fences(raw)


def call_prompt_1b(client: OpenAI, model: str, doc_text: str, existing_ttl: str) -> str:
    poml_path = PROMPTS_DIR / "refine_ontology.poml"
    with open(poml_path, "r", encoding="utf-8", errors="replace") as f:
        poml_content = f.read()
    
    params = poml(
        poml_content,
        context={"documents": doc_text, "existing_ttl": existing_ttl},
        format="openai_chat",
    )
    params["model"] = model

    response = client.chat.completions.create(**params)
    return strip_code_fences(response.choices[0].message.content)


def call_prompt_2(client: OpenAI, model: str, ttl_content: str, errors: str = "") -> str:
    poml_path = PROMPTS_DIR / "validate_syntax.poml"
    with open(poml_path, "r", encoding="utf-8", errors="replace") as f:
        poml_content = f.read()
    
    params = poml(
        poml_content,
        context={"ttl_content": ttl_content, "errors": errors},
        format="openai_chat",
    )
    params["model"] = model

    response = client.chat.completions.create(**params)
    return strip_code_fences(response.choices[0].message.content)


# ---------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------
def build_document_block(docs: list[dict]) -> str:
    parts = []
    for i, doc in enumerate(docs, 1):
        parts.append(f"--- Document {i}: {doc['name']} ---\n{doc['text']}")
    return "\n\n".join(parts)


def run_pipeline(uploaded_files, gen_model: str, val_model: str, refine: bool, status):
    client = OpenAI(api_key=OPENAI_API_KEY)
    doc_budget = KNOWN_MODEL_LIMITS[gen_model]["doc_budget"]

    # ── Extract text ──
    status.update(label="Extracting text from documents...")
    docs = []
    for f in uploaded_files:
        text = extract_text(f)
        if text.strip():
            tokens = count_tokens(text)
            docs.append({"name": f.name, "text": text, "tokens": tokens})

    if not docs:
        st.error("No text could be extracted from the uploaded files.")
        return None, []

    total_tokens = sum(d["tokens"] for d in docs)
    st.info(
        f"Loaded **{len(docs)} documents** — "
        f"**{total_tokens:,} tokens** total | "
        f"Model budget: **{doc_budget:,} tokens**"
    )

    # ── Step 1a: Generate initial ontology ──
    if total_tokens <= doc_budget:
        status.update(label=f"Step 1a: Generating ontology ({gen_model}) — direct path...")
        doc_text = build_document_block(docs)
        ttl_string = call_prompt_1a(client, gen_model, doc_text)
    else:
        status.update(label="Documents exceed budget. Chunking...")
        chunks = chunk_all_documents(docs)
        batches = batch_chunks(chunks)
        st.info(f"Split into **{len(chunks)} chunks** across **{len(batches)} batches**.")

        registry: set[str] = set()
        fragments: list[str] = []
        progress = st.progress(0, text="Processing batches...")

        for i, batch in enumerate(batches):
            status.update(label=f"Step 1a: Batch {i + 1}/{len(batches)} ({gen_model})...")
            batch_text = "\n\n---\n\n".join(batch)
            registry_text = ", ".join(sorted(registry)) if registry else ""

            raw_ttl = call_prompt_1a(client, gen_model, batch_text, registry_text)
            is_valid, result = validate_ttl(raw_ttl)

            if is_valid:
                fragments.append(result)
                new_names = extract_registry(result)
                registry.update(new_names)
                logger.info(
                    f"Batch {i + 1}: +{len(new_names)} names, "
                    f"registry total: {len(registry)}"
                )
            else:
                fragments.append(raw_ttl)
                logger.warning(f"Batch {i + 1} produced invalid TTL: {result}")

            progress.progress((i + 1) / len(batches), text=f"Batch {i + 1}/{len(batches)} done — registry: {len(registry)} names")

        status.update(label="Merging ontology fragments...")
        ttl_string = merge_ttl_fragments(fragments)
        st.success(f"Merged {len(fragments)} fragments. Registry: {len(registry)} names.")

    is_valid, result = validate_ttl(ttl_string)
    if is_valid:
        ttl_string = result
        st.success("Step 1a complete: Initial ontology generated.")
    else:
        st.warning("Step 1a: Initial TTL has syntax issues — will attempt to fix.")
        logger.warning(f"Step 1a TTL error: {result}")

    # ── Step 1b: Refine ontology (optional, same model) ──
    if refine:
        status.update(label=f"Step 1b: Refining ontology ({gen_model})...")
        doc_text = build_document_block(docs)
        additions = call_prompt_1b(client, gen_model, doc_text, ttl_string)

        if additions.strip().upper() == "NO_ADDITIONS_NEEDED":
            st.info("Step 1b: No additional types needed — ontology is comprehensive.")
        else:
            additions_valid, additions_result = validate_ttl(additions)
            if additions_valid:
                try:
                    ttl_string = merge_ttl_fragments([ttl_string, additions_result])
                    st.success("Step 1b complete: Merged new entity/relationship types.")
                except Exception as e:
                    logger.warning(f"Merge failed: {e}. Skipping refinement additions.")
                    st.warning("Step 1b: Could not merge additions. Proceeding with initial ontology.")
            else:
                logger.warning(f"Refinement additions had invalid TTL: {additions_result}")
                st.warning("Step 1b: Refinement output had syntax issues. Proceeding with initial ontology.")
    else:
        st.info("Step 1b: Refinement skipped.")

    # ── Step 2: Validate syntax (rdflib first, LLM only if needed) ──
    status.update(label="Step 2: Validating syntax (rdflib)...")
    is_valid, parse_result = validate_ttl(ttl_string)

    if is_valid:
        ttl_string = parse_result
        st.success("Step 2 complete: Syntax validated by rdflib — no LLM call needed.")
    else:
        max_attempts = 5
        current_ttl = ttl_string
        errors = parse_result

        for attempt in range(1, max_attempts + 1):
            st.warning(f"Step 2: Attempt {attempt}/{max_attempts} — calling {val_model} to fix syntax errors...")
            status.update(label=f"Step 2: Fix attempt {attempt}/{max_attempts} ({val_model})...")
            fixed_ttl = call_prompt_2(client, val_model, current_ttl, errors=errors)
            is_fixed, result = validate_ttl(fixed_ttl)

            if is_fixed:
                ttl_string = result
                st.success(f"Step 2 complete: Syntax fixed on attempt {attempt}.")
                break
            else:
                current_ttl = fixed_ttl
                errors = result

            if attempt == max_attempts:
                st.error(f"Step 2: Could not fix syntax errors after {max_attempts} attempts. Last error: {errors}")
                ttl_string = fixed_ttl

    status.update(label="Done!", state="complete")
    return ttl_string, docs


# ---------------------------------------------------------------------------
# Ontology stats
# ---------------------------------------------------------------------------
def get_ontology_stats(ttl_string: str) -> dict:
    g = Graph()
    g.parse(data=ttl_string, format="turtle")
    classes = set(g.subjects(RDF.type, OWL.Class))
    obj_props = set(g.subjects(RDF.type, OWL.ObjectProperty))
    data_props = set(g.subjects(RDF.type, OWL.DatatypeProperty))
    return {
        "Classes": len(classes),
        "Object Properties": len(obj_props),
        "Datatype Properties": len(data_props),
        "Total Triples": len(g),
    }


# ---------------------------------------------------------------------------
# Streamlit UI
# ---------------------------------------------------------------------------
def run_kg_pipeline(docs, ttl_string, gen_model, status, alpha=0.1):
    try:
        neo4j_client = Neo4jClient()
        driver = neo4j_client()
    except Exception as e:
        st.error(f"Could not connect to Neo4j: {e}")
        return

    try:
        # ── Snapshot before build ──
        stats_before = get_graph_stats(driver)

        # ── Convert ontology to schema ──
        status.update(label="Converting ontology to schema...")
        schema = ontology_to_schema(ttl_string)
        st.success(
            f"Schema: **{len(schema.node_types)}** node types, "
            f"**{len(schema.relationship_types)}** relationship types, "
            f"**{len(schema.patterns)}** patterns"
        )

        # ── Build KG ──
        status.update(label=f"Building Knowledge Graph ({gen_model})...")
        progress = st.progress(0, text="Processing documents...")

        def on_doc_complete(done, total, name):
            progress.progress(done / total, text=f"Document {done}/{total}: {name}")

        build_knowledge_graph(driver, schema, docs, gen_model, on_complete=on_doc_complete)
        st.success("Entity extraction and KG construction complete.")

        # ── Compute edge weights ──
        status.update(label="Computing edge weights...")
        weight_stats = compute_edge_weights(driver, alpha=alpha)
        st.success(f"Edge weights computed: **{weight_stats['updated']}** relationships (max shared chunks: {weight_stats['max_shared']}, α={alpha})")

        # ── Snapshot after build ──
        stats_after = get_graph_stats(driver)
        new_entities = stats_after["entities"] - stats_before["entities"]
        new_rels = stats_after["relationships"] - stats_before["relationships"]

        # ── Detect duplicates ──
        status.update(label="Detecting duplicate entities...")
        duplicates = find_duplicate_entities(driver)

        # ── Store results for dashboard ──
        st.session_state["kg_stats"] = {
            "before": stats_before,
            "after": stats_after,
            "new_entities": new_entities,
            "new_relationships": new_rels,
            "duplicates": duplicates,
        }

        status.update(label="Done!", state="complete")

    finally:
        neo4j_client.close()


def main():
    st.set_page_config(page_title="Ontology-Driven KG Builder", layout="wide")
    st.title("Ontology-Driven Knowledge Graph Builder")
    st.caption("Upload documents → Generate ontology → Build Knowledge Graph")

    # ── Sidebar ──
    model_names = list(KNOWN_MODEL_LIMITS.keys())
    with st.sidebar:
        st.header("Configuration")
        gen_model = st.selectbox("Generation Model", options=model_names, index=0)
        val_model = st.selectbox("Syntax Fix Model", options=model_names, index=1,
                                 help="Only used if rdflib detects syntax errors")
        st.divider()
        refine = st.checkbox("Refine ontology (Step 1b)", value=False,
                             help="Extra LLM pass to add missing entity types and relationships")
        budget = KNOWN_MODEL_LIMITS[gen_model]["doc_budget"]
        st.metric("Document Budget", f"{budget:,} tokens")
        st.divider()
        alpha = st.slider("Edge Weight Floor (α)", min_value=0.0, max_value=0.5, value=0.1, step=0.05,
                          help="Minimum weight for any relationship. Higher = less contrast between weak and strong edges.")
        st.divider()
        st.header("Knowledge Graph"
        )

    # ── File uploader ──
    uploaded_files = st.file_uploader(
        "Upload Documents",
        type=ACCEPTED_FILE_TYPES,
        accept_multiple_files=True,
        help="Supported formats: .txt, .pdf, .docx — upload related documents only",
    )

    if uploaded_files:
        file_info = []
        for f in uploaded_files:
            content = f.read()
            f.seek(0)
            size_kb = len(content) / 1024
            file_info.append({"File": f.name, "Size": f"{size_kb:.1f} KB"})
        st.dataframe(file_info, use_container_width=True, hide_index=True)

    # ── Step 1: Generate Ontology ──
    st.header("Step 1: Generate Ontology")

    if st.button("Generate Ontology", type="primary", disabled=not uploaded_files):
        with st.status("Starting ontology pipeline...", expanded=True) as status:
            ttl_result, docs = run_pipeline(uploaded_files, gen_model, val_model, refine, status)

        if ttl_result:
            st.session_state["ttl_result"] = ttl_result
            st.session_state["docs"] = docs

    ttl_result = st.session_state.get("ttl_result")
    docs = st.session_state.get("docs")

    if ttl_result:
        st.subheader("Ontology Summary")
        try:
            stats = get_ontology_stats(ttl_result)
            cols = st.columns(len(stats))
            for col, (label, value) in zip(cols, stats.items()):
                col.metric(label, value)
        except Exception:
            pass

        with st.expander("View Generated Ontology (.ttl)"):
            st.code(ttl_result, language="turtle")

        st.download_button(
            label="Download ontology.ttl",
            data=ttl_result,
            file_name="ontology.ttl",
            mime="text/turtle",
        )

        # ── Step 2: Build Knowledge Graph ──
        st.header("Step 2: Build Knowledge Graph")

        if st.button("Build Knowledge Graph", type="primary"):
            with st.status("Building Knowledge Graph...", expanded=True) as kg_status:
                run_kg_pipeline(docs, ttl_result, gen_model, kg_status, alpha=alpha)

    # ── Overlap Dashboard ──
    kg_stats = st.session_state.get("kg_stats")
    if kg_stats:
        st.divider()
        st.header("Overlap Dashboard")

        col1, col2, col3, col4 = st.columns(4)
        col1.metric("Entities Before", kg_stats["before"]["entities"])
        col2.metric("New Entities", kg_stats["new_entities"])
        col3.metric("Total Entities", kg_stats["after"]["entities"])
        col4.metric("Duplicate Pairs", len(kg_stats["duplicates"]))

        duplicates = kg_stats["duplicates"]
        if duplicates:
            same_label = [d for d in duplicates if d["match_type"] == "same_label"]
            cross_label = [d for d in duplicates if d["match_type"] == "cross_label"]
            overlap_pct = (len(duplicates) / kg_stats["after"]["entities"]) * 100 if kg_stats["after"]["entities"] else 0
            st.warning(f"**{len(duplicates)} duplicate entity pairs found** ({overlap_pct:.1f}% of total entities)")

            if same_label:
                with st.expander(f"Same-type duplicates — high confidence ({len(same_label)})"):
                    dup_display = [{"Name": d["name"], "Type": d["label_a"]} for d in same_label]
                    st.dataframe(dup_display, use_container_width=True, hide_index=True)

            if cross_label:
                with st.expander(f"Cross-type duplicates — needs review ({len(cross_label)})"):
                    dup_display = [{"Name": d["name"], "Type A": d["label_a"], "Type B": d["label_b"]} for d in cross_label]
                    st.dataframe(dup_display, use_container_width=True, hide_index=True)
        else:
            st.success("No duplicate entities found. Graph is clean.")

    # ── Entity Resolution ──
    run_entity_resolution_section()

    # ── Advanced Vector Embeddings ──
    run_advanced_embeddings_section()

    # ── External Web Sources ──
    run_web_sources_section(gen_model)

    # ── Graph Visualization ──
    st.divider()
    run_graph_visualization()

    # ── Step 3: Ask Questions ──
    st.divider()
    run_chat_section(gen_model)


CONFIDENCE_COLORS = {"high": "#2ecc71", "medium": "#f39c12", "low": "#e74c3c", "error": "#95a5a6"}


def run_entity_resolution_section():
    st.divider()
    st.header("Entity Resolution")

    try:
        neo4j_client = Neo4jClient()
        driver = neo4j_client()
        if not has_any_entities(driver):
            st.info("No entities in graph. Build a Knowledge Graph first.")
            neo4j_client.close()
            return
    except Exception:
        st.warning("Could not connect to Neo4j.")
        return

    # ── Scoring controls ──
    st.subheader("1. Find & Score Candidates")
    score_col1, score_col2 = st.columns(2)
    with score_col1:
        scoring_level = st.radio(
            "Scoring level",
            ["Exact match (free)", "Exact + Embedding ($)", "Exact + Embedding + LLM ($$)"],
            index=0,
            help="Higher levels cost more but catch fuzzy duplicates",
        )
    with score_col2:
        llm_model = st.selectbox("LLM Judge Model", ["gpt-4o-mini", "gpt-4o"], index=0,
                                  help="Only used at LLM scoring level")

    if st.button("Run Entity Resolution Scoring", type="primary"):
        with st.spinner("Finding candidate pairs..."):
            candidates = find_candidate_pairs(driver)
            all_pairs = candidates["same_label"] + candidates["cross_label"]

        if not all_pairs:
            st.success("No duplicate candidates found. Graph is clean.")
            neo4j_client.close()
            return

        st.info(f"Found **{len(candidates['same_label'])}** same-type and **{len(candidates['cross_label'])}** cross-type candidate pairs.")

        # Level 1: Exact
        scored = [score_exact(p) for p in all_pairs]

        # Level 2: Embedding
        if scoring_level in ("Exact + Embedding ($)", "Exact + Embedding + LLM ($$)"):
            with st.spinner("Computing embedding similarity..."):
                scored = score_embedding_batch(all_pairs)

        # Level 3: LLM judge (only for ambiguous pairs)
        if scoring_level == "Exact + Embedding + LLM ($$)":
            ambiguous = [s for s in scored if s["confidence"] in ("medium", "low")]
            if ambiguous:
                with st.spinner(f"LLM judging {len(ambiguous)} ambiguous pairs..."):
                    llm_scored = score_llm_batch(ambiguous, model=llm_model)
                    llm_map = {(s["id_a"], s["id_b"]): s for s in llm_scored}
                    scored = [llm_map.get((s["id_a"], s["id_b"]), s) for s in scored]

        st.session_state["er_scored_pairs"] = scored

    # ── Per-pair approval ──
    scored_pairs = st.session_state.get("er_scored_pairs")
    if scored_pairs:
        st.subheader("2. Review & Approve Pairs")
        st.caption("Check the pairs you want to merge. Uncheck to keep separate.")

        sorted_pairs = sorted(scored_pairs, key=lambda x: (
            {"high": 0, "medium": 1, "low": 2, "error": 3}.get(x["confidence"], 3),
            -x["score"],
        ))

        if "er_approvals" not in st.session_state:
            st.session_state["er_approvals"] = {
                i: p["confidence"] == "high" for i, p in enumerate(sorted_pairs)
            }

        for i, pair in enumerate(sorted_pairs):
            color = CONFIDENCE_COLORS.get(pair["confidence"], "#aaa")
            col_check, col_info, col_detail = st.columns([0.5, 3, 4])

            with col_check:
                st.session_state["er_approvals"][i] = st.checkbox(
                    "Merge", value=st.session_state["er_approvals"].get(i, False),
                    key=f"er_pair_{i}", label_visibility="collapsed",
                )

            with col_info:
                st.markdown(
                    f'<span style="color:{color}; font-weight:bold;">{pair["confidence"].upper()}</span> '
                    f'({pair["score_type"]}: {pair["score"]:.2f})  **{pair["name"]}**',
                    unsafe_allow_html=True,
                )

            with col_detail:
                type_info = (f'{pair["label_a"]}' if pair["label_a"] == pair["label_b"]
                             else f'{pair["label_a"]} vs {pair["label_b"]}')
                st.caption(f'{type_info} — {pair["reason"]}')

        approved = [sorted_pairs[i] for i, checked in st.session_state["er_approvals"].items() if checked]
        st.info(f"**{len(approved)}** of **{len(sorted_pairs)}** pairs approved for merge.")

        # ── Clustering + Merge ──
        st.subheader("3. Merge")

        if not approved:
            st.warning("No pairs approved. Check at least one pair above to enable merge.")
        else:
            clusters = build_transitive_clusters(approved)
            oversized = [c for c in clusters if len(c) > 5]

            with st.expander(f"Preview: {len(clusters)} merge clusters"):
                for ci, cluster in enumerate(clusters):
                    names = [f'{n["name"]} ({n["label"]})' for n in cluster]
                    survivor = names[0]
                    icon = "!" if len(cluster) > 5 else str(ci + 1)
                    st.markdown(f"**Cluster {icon}:** Keep **{survivor}**, merge: {', '.join(names[1:])}")

            if oversized:
                st.warning(f"{len(oversized)} cluster(s) exceed 5 members — review carefully before merging.")

            if st.button("Execute Merge", type="primary"):
                merge_log = []
                progress = st.progress(0, text="Merging clusters...")

                for ci, cluster in enumerate(clusters):
                    result = merge_cluster(driver, cluster)
                    merge_log.append(result)
                    progress.progress((ci + 1) / len(clusters), text=f"Merged cluster {ci + 1}/{len(clusters)}")

                st.session_state["merge_log"] = merge_log

                total_merged = sum(r["dropped_count"] for r in merge_log if r["status"] == "merged")
                st.success(f"Merged {total_merged} duplicate nodes across {len(clusters)} clusters.")

                # Refresh stats
                stats_after = get_graph_stats(driver)
                st.session_state["kg_stats"]["after"] = stats_after
                st.session_state["kg_stats"]["duplicates"] = find_duplicate_entities(driver)

                # Check for orphans
                orphans = find_orphaned_nodes(driver)
                if orphans:
                    st.warning(f"{len(orphans)} orphaned node(s) detected after merge:")
                    st.dataframe(
                        [{"Name": o["name"], "Type": o["label"]} for o in orphans],
                        use_container_width=True, hide_index=True,
                    )

                # Clear scored pairs since graph changed
                st.session_state.pop("er_scored_pairs", None)
                st.session_state.pop("er_approvals", None)

    # ── Undo ──
    merge_log = st.session_state.get("merge_log")
    if merge_log:
        all_snapshots = []
        for entry in merge_log:
            if entry.get("snapshots"):
                all_snapshots.extend(entry["snapshots"])

        if all_snapshots:
            st.divider()
            if st.button("Undo Last Merge", type="secondary"):
                with st.spinner("Restoring merged nodes..."):
                    restored = undo_merge(driver, all_snapshots)
                    st.success(f"Restored {restored} node(s). Note: re-run scoring to verify graph state.")
                    st.session_state.pop("merge_log", None)

    neo4j_client.close()

def run_advanced_embeddings_section():
    st.divider()
    st.header("Advanced Vector Embeddings")
    st.caption("Multi-index embeddings for entities and relationships beyond Chunk nodes")

    try:
        neo4j_client = Neo4jClient()
        driver = neo4j_client()
        if not has_any_entities(driver):
            st.info("No entities in graph. Build a Knowledge Graph first.")
            neo4j_client.close()
            return
    except Exception:
        st.warning("Could not connect to Neo4j.")
        return

    # Embedding Stats
    st.subheader("Embedding Statistics")
    if st.button("Refresh Embedding Stats"):
        with st.spinner("Computing embedding statistics..."):
            stats = get_embedding_stats(driver)
            st.session_state["embedding_stats"] = stats

    stats = st.session_state.get("embedding_stats")
    if stats:
        col1, col2, col3 = st.columns(3)
        col1.metric("Chunks with Embeddings", stats.get("chunks_with_embeddings", 0))
        col2.metric("Entities with Embeddings", sum(stats.get("entities_with_embeddings", {}).values()))
        col3.metric("Relationships with Embeddings", stats.get("relationships_with_embeddings", 0))

        if stats.get("entities_with_embeddings"):
            st.write("**Entities by type:**")
            for label, cnt in stats["entities_with_embeddings"].items():
                st.write(f"  - {label}: {cnt}")

    # Embed All Entity Types
    st.subheader("Embed All Entity Types")
    st.caption("Automatically detect entity types and create vector indexes for each")

    if st.button("Embed All Entity Types", type="primary"):
        with st.spinner("Embedding all entity types..."):
            results = embed_all_entity_types(driver)
            
        st.success(f"Embedded {sum(r['indexed'] for r in results.values())} entities across {len(results)} types")
        for entity_type, result in results.items():
            st.write(f"  - {entity_type}: {result['indexed']} indexed, properties: {result.get('properties', [])}")

    # ── Embed Specific Entity Type ──
    st.subheader("Embed Specific Entity Type")
    
    # Get available entity types
    with driver.session() as session:
        types = list(session.run("""
            MATCH (n) WHERE n.name IS NOT NULL 
              AND NONE(lbl IN labels(n) WHERE lbl IN ['Document', 'Chunk', 'WebDocument', 'WebChunk'])
            UNWIND labels(n) AS label DISTINCT label
        """))
        available_types = [dict(t)["label"] for t in types]

    if available_types:
        selected_type = st.selectbox("Select Entity Type", available_types)
        
        # Get properties for this type
        with driver.session() as session:
            props_result = list(session.run(f"""
                MATCH (n:{selected_type}) 
                WHERE n.name IS NOT NULL
                RETURN keys(n) AS keys
                LIMIT 1
            """))
            available_props = list(props_result[0].values())[0] if props_result else ["name"]
        
        embedding_props = st.multiselect(
            "Properties to include in embedding",
            options=available_props,
            default=["name", "description"] if "description" in available_props else ["name"],
        )
        
        if st.button(f"Embed {selected_type} Entities"):
            config = ENTITY_INDEX_CONFIGS.get(selected_type, {
                "index_name": f"{selected_type.lower()}_embedding_index",
                "dimensions": EMBEDDING_DIMENSIONS,
            })
            
            with st.spinner(f"Embedding {selected_type} entities..."):
                result = embed_entities_by_type(
                    driver,
                    entity_type=selected_type,
                    embedding_props=embedding_props,
                    index_name=config["index_name"],
                    dimensions=config.get("dimensions", EMBEDDING_DIMENSIONS),
                )
            
            st.success(f"Indexed {result['indexed']} {selected_type} entities")

    # Embed Relationships
    st.subheader("Embed Relationships")
    st.caption("Create embeddings for relationships based on source-type-target context")

    if st.button("Embed All Relationships", type="secondary"):
        with st.spinner("Embedding relationships..."):
            result = embed_relationships(driver)
        st.success(f"Indexed {result['indexed']} relationships")

    neo4j_client.close()


def run_web_sources_section(gen_model: str):
    st.divider()
    st.header("External Web Sources")
    st.caption("Augment your knowledge graph with related web content")

    ttl_result = st.session_state.get("ttl_result")
    docs = st.session_state.get("docs")
    if not ttl_result or not docs:
        st.info("Generate an ontology and build a Knowledge Graph first.")
        return

    try:
        neo4j_client = Neo4jClient()
        driver = neo4j_client()
        if not has_any_entities(driver):
            st.info("No entities in graph. Build a Knowledge Graph first.")
            neo4j_client.close()
            return
    except Exception:
        st.warning("Could not connect to Neo4j.")
        return

    # Show existing web stats if present
    try:
        ws = get_web_source_stats(driver)
        if ws["web_chunks"] > 0:
            st.session_state["web_sources_enabled"] = True
            st.success(f"Web content active: **{ws['web_chunks']}** web chunks, **{ws['similar_to_edges']}** SIMILAR_TO edges (avg weight: {ws['avg_edge_weight']:.3f})")
    except Exception:
        pass

    col1, col2, col3 = st.columns(3)
    with col1:
        max_articles = st.slider("Max articles", 1, 10, 5, key="ws_max_articles")
    with col2:
        sim_threshold = st.slider("Similarity threshold", 0.3, 0.8, 0.5, 0.05,
                                  key="ws_sim_threshold",
                                  help="Min cosine similarity to create SIMILAR_TO edge")
    with col3:
        max_topics = st.slider("Max topics", 2, 7, 5, key="ws_max_topics")

    if st.button("Fetch & Process Web Sources", type="primary"):
        with st.status("Processing web sources...", expanded=True) as ws_status:
            ws_status.update(label="Extracting key topics from your documents...")
            topics = extract_topics(docs, model=gen_model, max_topics=max_topics)
            st.info(f"Topics: {', '.join(topics)}")

            ws_status.update(label="Searching the web via OpenAI...")
            web_docs = search_and_fetch(topics[:max_articles], model=gen_model)

            if not web_docs:
                st.warning("No content found. Try different documents or topics.")
                ws_status.update(label="No results", state="error")
                neo4j_client.close()
                return

            st.info(f"Fetched web content for **{len(web_docs)}** topics")
            with st.expander("Web search results"):
                for r in web_docs:
                    st.markdown(f"- **{r['name']}** — {len(r['text'])} chars")

            ws_status.update(label=f"Building web knowledge graph ({gen_model})...")
            schema = ontology_to_schema(ttl_result)
            progress = st.progress(0, text="Processing web articles...")

            def on_web_complete(done, total, name):
                progress.progress(done / total, text=f"Web article {done}/{total}: {name}")

            build_web_knowledge_graph(driver, schema, web_docs, gen_model, on_complete=on_web_complete)

            ws_status.update(label="Computing SIMILAR_TO edges...")
            edge_stats = compute_similar_to_edges(driver, similarity_threshold=sim_threshold)

            st.session_state["web_sources_enabled"] = True
            st.session_state["web_source_stats"] = edge_stats

            ws_status.update(label="Done!", state="complete")

    ws_stats = st.session_state.get("web_source_stats")
    if ws_stats:
        mcol1, mcol2, mcol3 = st.columns(3)
        mcol1.metric("SIMILAR_TO Edges", ws_stats["edges_created"])
        mcol2.metric("Avg Similarity", f"{ws_stats['avg_similarity']:.3f}")
        mcol3.metric("Max Similarity", f"{ws_stats['max_similarity']:.3f}")

    if st.session_state.get("web_sources_enabled"):
        if st.button("Remove All Web Content", type="secondary"):
            with st.spinner("Removing web content..."):
                result = remove_web_content(driver)
                st.success(
                    f"Removed {result['web_chunks_removed']} web chunks, "
                    f"{result['web_documents_removed']} web documents, "
                    f"{result['edges_removed']} SIMILAR_TO edges."
                )
                st.session_state["web_sources_enabled"] = False
                st.session_state.pop("web_source_stats", None)

    neo4j_client.close()


COLOR_PALETTE = [
    "#4e79a7", "#f28e2b", "#e15759", "#76b7b2", "#59a14f",
    "#edc948", "#b07aa1", "#ff9da7", "#9c755f", "#bab0ac",
    "#86bcb6", "#8cd17d", "#a0cbe8", "#d4a6c8", "#ffbe7d",
    "#d7b5a6", "#b6992d", "#ff6b6b", "#4ecdc4", "#45b7d1",
]


def run_graph_visualization():
    st.header("Graph Visualization")

    try:
        neo4j_client = Neo4jClient()
        driver = neo4j_client()
        if not has_any_entities(driver):
            st.info("No graph to visualize. Build a Knowledge Graph first.")
            neo4j_client.close()
            return
    except Exception:
        st.warning("Could not connect to Neo4j.")
        return

    node_limit = st.slider("Max nodes", min_value=20, max_value=500, value=100, step=20, key="viz_limit")

    if st.button("Visualize Graph"):
        with st.spinner("Loading graph..."):
            try:
                html, legend = build_graph_html(driver, node_limit)
                if legend:
                    cols = st.columns(min(len(legend), 6))
                    for i, (label, color) in enumerate(legend.items()):
                        cols[i % len(cols)].markdown(
                            f'<span style="color:{color}; font-size:20px;">&#9679;</span> {label}',
                            unsafe_allow_html=True,
                        )
                components.html(html, height=700, scrolling=False)
            except Exception as e:
                st.error(f"Visualization error: {e}")

    neo4j_client.close()


def build_graph_html(driver, limit: int) -> tuple[str, dict]:
    node_query = """
    MATCH (n)
    WHERE n.name IS NOT NULL
      AND NONE(lbl IN labels(n) WHERE lbl IN ['Document', 'Chunk', 'WebDocument', 'WebChunk'])
    RETURN elementId(n) AS id, n.name AS name, labels(n) AS labels
    LIMIT $limit
    """
    rel_query = """
    MATCH (a)-[r]->(b)
    WHERE a.name IS NOT NULL AND b.name IS NOT NULL
      AND NONE(lbl IN labels(a) WHERE lbl IN ['Document', 'Chunk', 'WebDocument', 'WebChunk'])
      AND NONE(lbl IN labels(b) WHERE lbl IN ['Document', 'Chunk', 'WebDocument', 'WebChunk'])
      AND NOT type(r) IN ['FROM_CHUNK', 'FROM_DOCUMENT', 'NEXT_CHUNK', 'SIMILAR_TO']
    WITH a, r, b LIMIT $limit
    RETURN elementId(a) AS source, elementId(b) AS target, type(r) AS rel_type,
           coalesce(r.weight, 1.0) AS weight
    """

    with driver.session() as session:
        nodes = session.run(node_query, limit=limit).data()
        rels = session.run(rel_query, limit=limit * 3).data()

    # Count connections per node for sizing
    degree: dict[str, int] = {}
    for rel in rels:
        degree[rel["source"]] = degree.get(rel["source"], 0) + 1
        degree[rel["target"]] = degree.get(rel["target"], 0) + 1

    # Auto-assign colors to entity types
    all_types = set()
    for node in nodes:
        for lbl in node["labels"]:
            if lbl not in ("__Entity__", "Document", "Chunk", "WebDocument", "WebChunk"):
                all_types.add(lbl)
    type_colors = {}
    for i, t in enumerate(sorted(all_types)):
        type_colors[t] = COLOR_PALETTE[i % len(COLOR_PALETTE)]

    net = Network(height="680px", width="100%", bgcolor="#0e1117", font_color="#ffffff")
    net.barnes_hut(gravity=-5000, central_gravity=0.35, spring_length=200, spring_strength=0.01)

    node_ids = set()
    for node in nodes:
        nid = node["id"]
        node_ids.add(nid)
        name = node["name"]
        node_labels = [l for l in node["labels"] if l not in ("__Entity__",)]
        entity_type = node_labels[0] if node_labels else "Unknown"
        color = type_colors.get(entity_type, "#aaaaaa")
        size = 12 + min(degree.get(nid, 0) * 4, 40)
        short_name = name if len(name) <= 25 else name[:22] + "..."

        net.add_node(
            nid,
            label=short_name,
            title=f"<b>{entity_type}</b><br>{name}<br>Connections: {degree.get(nid, 0)}",
            color=color,
            size=size,
            font={"size": 11, "color": "#ffffff", "strokeWidth": 2, "strokeColor": "#000000"},
        )

    for rel in rels:
        if rel["source"] in node_ids and rel["target"] in node_ids:
            w = rel.get("weight", 1.0)
            edge_width = 0.5 + w * 3.0
            opacity = max(0.3, w)
            net.add_edge(
                rel["source"],
                rel["target"],
                title=f'{rel["rel_type"]} (weight: {w:.2f})',
                label=rel["rel_type"],
                color={"color": f"rgba(85,85,85,{opacity})", "highlight": "#ffffff"},
                font={"size": 8, "color": "#aaaaaa", "strokeWidth": 0},
                arrows="to",
                width=edge_width,
            )

    return net.generate_html(), type_colors


def run_chat_section(gen_model: str):
    st.header("Step 3: Ask Questions (GraphRAG)")

    try:
        neo4j_client = Neo4jClient()
        driver = neo4j_client()
        has_data = has_any_entities(driver)
        neo4j_client.close()
    except Exception:
        st.warning("Could not connect to Neo4j. Build a Knowledge Graph first.")
        return

    if not has_data:
        st.info("No entities found in the graph. Build a Knowledge Graph first.")
        return

    ret_col1, ret_col2, ret_col3, ret_col4 = st.columns(4)
    with ret_col1:
        retrieval_mode = st.selectbox(
            "Retrieval mode",
            ["Auto", "Hybrid", "Hybrid Enriched", "Vector", "Graph", "Fuzzy"],
            index=0,
        )
    with ret_col2:
        hops = st.select_slider("Reasoning hops", options=[1, 2], value=2,
                                help="1 = direct relationships only. 2 = multi-hop reasoning (follows neighbor's neighbors).")
    with ret_col3:
        weight_threshold = st.slider("Weight threshold", min_value=0.0, max_value=0.5, value=0.1, step=0.05,
                                     help="Ignore edges below this weight during retrieval.")
    with ret_col4:
        web_available = st.session_state.get("web_sources_enabled", False)
        include_web = st.checkbox("Include web sources", value=False,
                                  disabled=not web_available,
                                  help="Follow SIMILAR_TO edges to include supplementary web context" if web_available
                                       else "Fetch web sources first to enable this")

    if "chat_history" not in st.session_state:
        st.session_state["chat_history"] = []

    for msg in st.session_state["chat_history"]:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])

    question = st.chat_input("Ask a question about your knowledge graph...")

    if question:
        st.session_state["chat_history"].append({"role": "user", "content": question})
        with st.chat_message("user"):
            st.markdown(question)

        with st.chat_message("assistant"):
            with st.spinner("Searching knowledge graph..."):
                try:
                    neo4j_client = Neo4jClient()
                    driver = neo4j_client()
                    result = query_agentic_rag(driver, question, gen_model,
                                               mode=retrieval_mode,
                                               hops=hops,
                                               weight_threshold=weight_threshold,
                                               top_k=5,
                                               include_web_sources=include_web)
                    neo4j_client.close()

                    answer = result["answer"]
                    st.markdown(answer)
                    st.session_state["chat_history"].append({"role": "assistant", "content": answer})

                    if result.get("selected_mode") or result.get("routing_reason"):
                        with st.expander("Retrieval route"):
                            st.write(f"**Selected mode:** {result.get('selected_mode', result.get('retriever', 'N/A'))}")
                            st.write(result.get("routing_reason", "No routing reason returned."))
                            if result.get("latency_ms") is not None:
                                st.write(f"**Latency:** {result['latency_ms']} ms")

                    context = result.get("context")
                    if context and hasattr(context, "items") and context.items:
                        with st.expander("Retrieved context"):
                            for i, item in enumerate(context.items, 1):
                                st.markdown(f"**Chunk {i}** (score: {item.metadata.get('score', 'N/A') if item.metadata else 'N/A'})")
                                st.text(str(item.content)[:500])
                                st.divider()
                    elif context and isinstance(context, list):
                        with st.expander("Retrieved context"):
                            for i, item in enumerate(context, 1):
                                st.markdown(f"**Context {i}**")
                                st.text(str(item)[:1000])
                                st.divider()

                except Exception as e:
                    error_msg = f"Error querying graph: {e}"
                    st.error(error_msg)
                    st.session_state["chat_history"].append({"role": "assistant", "content": error_msg})


if __name__ == "__main__":
    main()
